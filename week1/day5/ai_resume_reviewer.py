import os
import json
from pathlib import Path
from groq import Groq
from dotenv import load_dotenv
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document
import time
import re
import docx
print(docx.__file__)

def extract_json(text):
    match = re.search(r'\{[\s\S]*\}', text)
    if match:
        return match.group(0)
    return text

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API ERROR")

client = Groq(api_key=my_api_key)    
model = "allam-2-7b"

job_description = """
AI DEVELOPER INTERN
BUILD • LEARN • INNOVATE
We are looking for an enthusiastic AI Developer Intern who is interested in
building intelligent applications and exploring the rapidly evolving world of
Artificial Intelligence and Machine Learning. This internship provides an opportunity to work on practical AI projects, strengthen programming skills, and understand how AI can be applied to real- world business and technology challenges. YOUR MISSION
As an AI Developer Intern, you will work with the technical team to develop, test, and improve AI-powered solutions. You will be involved in:
 Developing AI and Machine Learning applications.  Writing and testing code primarily using Python.  Preparing, cleaning, and processing datasets.  Training and evaluating basic ML models.  Implementing classification, prediction, recommendation, and automation
solutions.  Exploring Generative AI and Large Language Models (LLMs).  Building basic AI-powered applications such as chatbots and intelligent
assistants.  Working with AI APIs and integrating models into applications.  Performing basic prompt engineering and evaluating AI outputs.  Working with data using tools such as NumPy, Pandas, and Scikit-learn.  Debugging code and troubleshooting model or application issues.  Researching new AI technologies, models, and development approaches.  Collaborating with developers and other team members on AI projects.

YOUR TECH TOOLKIT
Programming
Python | C++ | Java | JavaScript
AI / ML
Machine Learning | Deep Learning | NLP | Generative AI | LLMs
Libraries & Frameworks
NumPy | Pandas | Scikit-learn | TensorFlow | PyTorch
Development
REST APIs | JSON | Databases | Git | GitHub
Advantage
Knowledge of LangChain, Hugging Face, OpenAI-compatible APIs, computer
vision, or chatbot development will be an added advantage. You don't need to know everything. Strong fundamentals and a willingness to
learn are what matter most. WHO CAN APPLY?
This opportunity is suitable for:
 Students and recent graduates from Computer Science, Artificial
Intelligence, Data Science, Machine Learning, Information Technology, or related fields.  Freshers with a strong interest in AI development.  Candidates who enjoy coding, problem-solving, experimentation, and
technology research. We especially value candidates who have:
✓ AI/ML academic projects
✓ Personal AI applications
✓ Chatbot projects

✓ Python projects
✓ Machine Learning models
✓ GitHub repositories
✓ Hackathon or research experience

WHAT YOU'LL DEVELOP
TECHNICAL SKILLS
Strengthen Python, Machine Learning, data processing, and AI development
skills. REAL PROJECT EXPERIENCE
Work on practical AI applications and business-oriented use cases. PROBLEM-SOLVING
Learn how to approach real-world problems using AI technologies. AI EXPOSURE
Explore modern developments in Generative AI, LLMs, NLP, and Machine
Learning. PORTFOLIO
Build meaningful AI projects that can strengthen your career profile. CAREER OUTCOME
This internship is designed to help you move from learning AI → building AI
→ solving real-world problems with AI.
"""
class JobD(BaseModel):
    role: str
    required_skills: list[str] = []
    preferred_skills: list[str] = []
    experience: float | None = None
    responsibility: list[str] = []
    educational_requirement: list[str] = []

    model_config = {"extra": "ignore"}

jobd_schema = JobD.model_json_schema()

system_prompt = """
Extract job info as JSON with these fields only:
{
  "role": "job title",
  "required_skills": ["skill1", "skill2"],
  "preferred_skills": ["skill1"],
  "experience": 0.0 or null,
  "responsibility": ["duty1"],
  "educational_requirement": ["req1"]
}
Return ONLY valid JSON. No markdown, no extra text.
"""

user_prompt = f"""
Analyze the following job description:
{job_description}
"""
message_system = {
    "role":"system",
    "content": system_prompt
}

message_user = {
    "role":"user",
    "content":user_prompt
}

response_format = {
    "type":"json_object"
}

messages = [message_system, message_user]

response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)

answer = response.choices[0].message.content

raw_json = answer

job_data = json.loads(extract_json(raw_json))

job = JobD(**job_data)

print(job.experience)
print(job.educational_requirement)

class MatchResult (BaseModel):
    score : float
    candidate_name: str
    matching_skills: list[str]
    missing_skills: list[str]
    experience_met: bool 
    verdict: str

class Experience(BaseModel):
    company : str | None = None
    role : str | None = None
    duraction : str | None = None
    description : str | None = None
    skills_used : list[str] = []

class Resume(BaseModel):
    name : str | None = None
    email : str | None = None
    phone : str | None = None

    total_experience_years : float | None = None

    skills : list[str] = []
    experience : list[str] = []
    education : list[str] = []
    project : list[str] = []
    Certificate : list[str] = []

resume_schema = Resume.model_json_schema()


def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + '\n'
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + '\n'

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                text += cell.text + "\n"

    return text

def read_resume(file_path):
    if file_path.suffix.lower()==".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower()==".docx":
        return read_docx(file_path)
    else:
        return None

def parse_resume(file_path):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:
    {resume_text}
    """
    message_system = {
        "role":"system",
         "content": system_prompt
    }

    message_user = {
        "role":"user",
        "content": user_prompt
    }

    messages = [message_system, message_user]
    response_format = {
        "type":"json_object"
    }
    response = client.chat.completions.create(model=model, response_format=response_format, messages=messages)
    raw_output = response.choices[0].message.content
    data = json.loads(extract_json(raw_output))
    resume = Resume(**data)
    return resume

def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    response_format = {
        "type":"json_object"
    }
    message = {
        "role":"user",
        "content":prompt
    }
    messages=[message]
    response = client.chat.completions.create(model=model, response_format=response_format, messages=messages)
    raw_data = response.choices[0].message.content
    data = json.loads(extract_json(raw_data))
    return MatchResult(**data)

resume_folder = Path(r"C:\users\aditya\ai_engineer_course\week1\day5\resumes")
all_result = []
for file_path in resume_folder.iterdir(): ##
    if file_path.suffix.lower() not in [".pdf",".docx"]:
        continue
    print("\nProcessing",file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume = parse_resume(resume_text)
    time.sleep(5)
    result = final_score(job, parsed_resume)
    time.sleep(5)
    print("Score:", result.score)
    all_result.append({
        "name":parsed_resume.name,
        "score":result.score,
        "skills":result.matching_skills,
        "missing skills":result.missing_skills,
        "experience":result.experience_met,
        "verdict":result.verdict
    })

all_result.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)

top_2 = all_result[:2]
lowest_2 = all_result[-2:]

print("TOP 2 CANDIDATES")
for candidate in top_2:
    print(
        candidate["name"],"-",candidate["score"],"%"
    )
    print(candidate["verdict"])

print("LOWEST 2 CANDIDATES")
for candidate in lowest_2:
    print(
        candidate["name"],"-",candidate["score"],"%"
    )
    print(candidate["verdict"])    


    